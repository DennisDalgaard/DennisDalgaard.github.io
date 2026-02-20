#!/usr/bin/env python3
"""Generate Excel overview of all Shelly devices, use cases, and filters."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

# ─── Color scheme ────────────────────────────────────────────
BLUE_FILL = PatternFill(start_color="0066FF", end_color="0066FF", fill_type="solid")
LIGHT_BLUE_FILL = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
DARK_FILL = PatternFill(start_color="0D1B2A", end_color="0D1B2A", fill_type="solid")
GREEN_FILL = PatternFill(start_color="E8F8EF", end_color="E8F8EF", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFF9E6", end_color="FFF9E6", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
GRAY_FILL = PatternFill(start_color="F7F8FC", end_color="F7F8FC", fill_type="solid")

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
SUBHEADER_FONT = Font(name="Calibri", bold=True, color="0D1B2A", size=11)
NORMAL_FONT = Font(name="Calibri", color="1A1A2E", size=10)
LINK_FONT = Font(name="Calibri", color="0066FF", size=10, underline="single")

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E6EF"),
    right=Side(style="thin", color="E2E6EF"),
    top=Side(style="thin", color="E2E6EF"),
    bottom=Side(style="thin", color="E2E6EF"),
)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_header(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = BLUE_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def style_row(ws, row, cols, fill=None):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = NORMAL_FONT
        cell.alignment = LEFT
        cell.border = THIN_BORDER
        if fill:
            cell.fill = fill


# ══════════════════════════════════════════════════════════════
# Sheet 1: ENHEDER (Devices)
# ══════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Enheder"

headers = [
    "ID", "Navn", "Serie", "Beskrivelse",
    "Kanaler", "Max Strøm", "Spænding(er)",
    "Installation", "Effektmåling", "Tørkontakt",
    "Rullegardin", "Protokoller",
    "Use Cases (ID'er)", "Tags",
    "HANDLING (behold/fjern/opdater)", "NOTER"
]

for col, header in enumerate(headers, 1):
    ws1.cell(row=1, column=col, value=header)
style_header(ws1, 1, len(headers))

devices = [
    ["shelly-1-gen4", "Shelly 1 Gen4", "Gen4",
     "Kompakt 1-kanals relæ med tørkontakt. WiFi, BT, Zigbee, Matter.",
     1, "16A", "110-240V AC, 12V DC, 24V DC",
     "flush", "Nej", "Ja", "Nej",
     "WiFi, Bluetooth, Zigbee, Matter",
     "lighting, garage, appliance, heating, intercom, motor",
     "Tørkontakt, Matter-klar", "", ""],

    ["shelly-1pm-gen4", "Shelly 1PM Gen4", "Gen4",
     "1-kanals relæ med effektmåling. 35% mindre end forgængeren.",
     1, "16A", "110-240V AC, 24V DC",
     "flush", "Ja", "Nej", "Nej",
     "WiFi, Bluetooth, Zigbee, Matter",
     "lighting, appliance, heating, energy",
     "Effektmåling, Ultra-kompakt", "", ""],

    ["shelly-1pm-mini-gen4", "Shelly 1PM Mini Gen4", "Gen4",
     "Verdens mindste relæ med effektmåling.",
     1, "8A", "110-240V AC",
     "flush", "Ja", "Nej", "Nej",
     "WiFi, Bluetooth, Zigbee, Matter",
     "lighting, appliance",
     "Effektmåling, Ultrakompakt", "", ""],

    ["shelly-2pm-gen4", "Shelly 2PM Gen4", "Gen4",
     "2-kanals relæ med effektmåling. Til persienner og bi-dir. motorer.",
     2, "10A/kanal", "110-240V AC, 24V DC",
     "flush", "Ja", "Nej", "Ja",
     "WiFi, Bluetooth, Zigbee, Matter",
     "blinds, lighting, appliance, motor",
     "Rullegardin, Effektmåling", "", ""],

    ["shelly-mini-1-gen4", "Shelly Mini 1 Gen4", "Gen4",
     "Ultrakompakt relæ til lampeudtag og trange dåser.",
     1, "8A", "110-240V AC",
     "flush", "Nej", "Nej", "Nej",
     "WiFi, Bluetooth, Zigbee, Matter",
     "lighting, appliance",
     "Ultrakompakt, Lampeudtag", "", ""],

    ["shelly-em-mini-gen4", "Shelly EM Mini Gen4", "Gen4",
     "Kompakt energimåler til individuelle apparater.",
     1, "N/A", "110-240V AC",
     "flush", "Ja", "Nej", "Nej",
     "WiFi, Bluetooth, Zigbee, Matter",
     "energy",
     "Energimåling, Kompakt", "", ""],

    ["shelly-plug-gen4", "Shelly Plug Gen4", "Gen4",
     "Smart stik med effektmåling op til 1800W/15A. Lux-sensor.",
     1, "15A", "110-240V AC",
     "plug", "Ja", "Nej", "Nej",
     "WiFi, Bluetooth, Zigbee, Matter",
     "appliance, energy, lighting",
     "Plug & Play, Effektmåling", "", ""],

    ["shelly-flood-gen4", "Shelly Flood Sensor Gen4", "Gen4 Sensor",
     "Vandlækagesensor med kabel. Realtidsalarm.",
     0, "N/A", "Batteri",
     "standalone", "Nej", "Nej", "Nej",
     "WiFi, Bluetooth",
     "sensors",
     "Batteridrevet, Alarm", "", ""],

    ["shelly-motion-gen4", "Shelly Motion Sensor Gen4", "Gen4 Sensor",
     "Bevægelsessensor med tilstedeværelsesregistrering.",
     0, "N/A", "Batteri",
     "standalone", "Nej", "Nej", "Nej",
     "WiFi, Bluetooth",
     "sensors, lighting, heating",
     "Tilstedeværelse, Batteridrevet", "", ""],

    ["shelly-dimmer-gen3", "Shelly Dimmer Gen3", "Gen3",
     "Lysdæmper til LED, halogener og transformatorer. Med/uden nul.",
     1, "2A (dimmer)", "110-240V AC",
     "flush", "Ja", "Nej", "Nej",
     "WiFi, Bluetooth",
     "dimming, lighting",
     "Lysdæmpning, Uden nul", "", ""],

    ["shelly-dimmer-010v-gen3", "Shelly Dimmer 0/1-10V PM Gen3", "Gen3",
     "0-10V/1-10V dæmper til LED-drivere, motorer og ventiler.",
     1, "2A", "110-240V AC",
     "flush", "Ja", "Nej", "Nej",
     "WiFi, Bluetooth",
     "dimming, lighting, motor",
     "0-10V, Effektmåling", "", ""],

    ["shelly-plus-1", "Shelly Plus 1", "Plus",
     "1-kanals relæ med tørkontakt (potentialfri). AC og DC forsyning.",
     1, "16A", "110-240V AC, 12V DC, 24-240V DC",
     "flush", "Nej", "Ja", "Nej",
     "WiFi, Bluetooth",
     "lighting, garage, appliance, intercom, heating",
     "Tørkontakt, DC-forsyning", "", ""],

    ["shelly-plus-1pm", "Shelly Plus 1PM", "Plus",
     "1-kanals relæ med effektmåling. 16A. AC og DC forsyning.",
     1, "16A", "110-240V AC, 24-30V DC",
     "flush", "Ja", "Nej", "Nej",
     "WiFi, Bluetooth",
     "lighting, appliance, heating, energy",
     "Effektmåling, 16A", "", ""],

    ["shelly-plus-2pm", "Shelly Plus 2PM", "Plus",
     "2-kanals relæ med effektmåling pr. kanal. Rullegardin-tilstand.",
     2, "10A/kanal", "110-240V AC, 24V DC",
     "flush", "Ja", "Nej", "Ja",
     "WiFi, Bluetooth",
     "blinds, lighting, motor",
     "Rullegardin, 2 kanaler", "", ""],

    ["shelly-plus-i4", "Shelly Plus i4", "Plus",
     "4-kanals scenecontroller uden relæ. 12 tilpasselige handlinger.",
     4, "N/A", "110-240V AC",
     "flush", "Nej", "Nej", "Nej",
     "WiFi, Bluetooth",
     "scenes",
     "Scener, 4 inputs", "", ""],

    ["shelly-plus-i4-dc", "Shelly Plus i4 DC", "Plus",
     "4-kanals scenecontroller til DC (5-24V).",
     4, "N/A", "5-24V DC",
     "flush", "Nej", "Nej", "Nej",
     "WiFi, Bluetooth",
     "scenes, sensors",
     "DC, Scener", "", ""],

    ["shelly-plus-uni", "Shelly Plus Uni", "Plus",
     "Universelt lavspændingsmodul. 2 SSR, analog/digital input, 1-Wire.",
     2, "0.3A", "8-24V AC, 5-28V DC",
     "flush", "Nej", "Ja", "Nej",
     "WiFi, Bluetooth",
     "intercom, sensors, garage",
     "Lavspænding, Tørkontakt, Sensorinput", "", ""],

    ["shelly-pro-1", "Shelly Pro 1", "Pro",
     "Professionelt 1-kanals DIN-skinne relæ med tørkontakt.",
     1, "16A", "110-240V AC",
     "din", "Nej", "Ja", "Nej",
     "WiFi, LAN, Bluetooth",
     "lighting, garage, appliance, motor, heating",
     "DIN-skinne, LAN", "", ""],

    ["shelly-pro-1pm", "Shelly Pro 1PM", "Pro",
     "Professionelt 1-kanals DIN-skinne relæ med effektmåling.",
     1, "16A", "110-240V AC",
     "din", "Ja", "Nej", "Nej",
     "WiFi, LAN, Bluetooth",
     "lighting, appliance, heating, energy",
     "DIN-skinne, LAN, Effektmåling", "", ""],

    ["shelly-pro-2pm", "Shelly Pro 2PM", "Pro",
     "2-kanals DIN-skinne relæ. 16A/kanal, 25A total. Rullegardin.",
     2, "16A/kanal (25A total)", "110-240V AC",
     "din", "Ja", "Nej", "Ja",
     "WiFi, LAN, Bluetooth",
     "blinds, lighting, motor",
     "DIN-skinne, Rullegardin, LAN", "", ""],

    ["shelly-pro-4pm", "Shelly Pro 4PM", "Pro",
     "4-kanals DIN-skinne relæ med farveskærm. 16A/kanal, 40A total.",
     4, "16A/kanal (40A total)", "110-240V AC",
     "din", "Ja", "Nej", "Nej",
     "WiFi, LAN, Bluetooth",
     "lighting, appliance, heating, energy",
     "DIN-skinne, 4 kanaler, Farveskærm", "", ""],

    ["shelly-pro-dimmer-1pm", "Shelly Pro Dimmer 1PM", "Pro",
     "1-kanals DIN-skinne dæmper med effektmåling.",
     1, "10A", "110-240V AC",
     "din", "Ja", "Nej", "Nej",
     "WiFi, LAN, Bluetooth",
     "dimming, lighting",
     "DIN-skinne, Lysdæmpning", "", ""],

    ["shelly-pro-dimmer-2pm", "Shelly Pro Dimmer 2PM", "Pro",
     "2-kanals DIN-skinne dæmper med effektmåling pr. kanal.",
     2, "10A", "110-240V AC",
     "din", "Ja", "Nej", "Nej",
     "WiFi, LAN, Bluetooth",
     "dimming, lighting",
     "DIN-skinne, Lysdæmpning, 2 kanaler", "", ""],

    ["shelly-pro-dimmer-010v", "Shelly Pro Dimmer 0/1-10V PM", "Pro",
     "DIN-skinne 0-10V/1-10V dæmper til LED-drivere og motorer.",
     1, "N/A", "110-240V AC",
     "din", "Ja", "Nej", "Nej",
     "WiFi, LAN, Bluetooth",
     "dimming, motor",
     "DIN-skinne, 0-10V, LAN", "", ""],

    ["shelly-pro-3em", "Shelly Pro 3EM", "Pro",
     "3-faset DIN-skinne energimåler med 120A CT-spoler. MODBUS.",
     3, "120A CT", "110-240V AC",
     "din", "Ja", "Nej", "Nej",
     "WiFi, LAN, Bluetooth",
     "energy",
     "DIN-skinne, 3-faset, MODBUS", "", ""],

    ["shelly-pro-3em-400", "Shelly Pro 3EM-400", "Pro",
     "3-faset DIN-skinne energimåler med 400A CT-spoler.",
     3, "400A CT", "110-240V AC",
     "din", "Ja", "Nej", "Nej",
     "WiFi, LAN, Bluetooth",
     "energy",
     "DIN-skinne, 3-faset, 400A", "", ""],

    ["shelly-wall-display", "Shelly Wall Display X2", "Display",
     "6,95\" touchskærm. Integreret relæ, effektmåling og sensorer.",
     2, "10A", "110-240V AC",
     "flush", "Ja", "Nej", "Nej",
     "WiFi, Bluetooth",
     "display, scenes, lighting",
     "Touchskærm, Sensorer", "", ""],

    ["shelly-wall-display-xl", "Shelly Wall Display XL", "Display",
     "10\" touchskærm kontrolpanel. Scener, energi og vejrdata.",
     0, "N/A", "110-240V AC",
     "flush", "Nej", "Nej", "Nej",
     "WiFi, Bluetooth",
     "display, scenes",
     "10\" skærm, Sonos-integration", "", ""],
]

for i, device in enumerate(devices):
    row = i + 2
    for col, val in enumerate(device, 1):
        ws1.cell(row=row, column=col, value=val)
    fill = GRAY_FILL if i % 2 == 0 else WHITE_FILL
    style_row(ws1, row, len(headers), fill)

# Last two columns (HANDLING, NOTER) get yellow fill for editing
for i in range(len(devices)):
    row = i + 2
    ws1.cell(row=row, column=15).fill = YELLOW_FILL
    ws1.cell(row=row, column=16).fill = YELLOW_FILL

# Add some empty rows for new devices
for i in range(10):
    row = len(devices) + 2 + i
    for col in range(1, len(headers) + 1):
        cell = ws1.cell(row=row, column=col)
        cell.fill = GREEN_FILL
        cell.border = THIN_BORDER
        cell.font = NORMAL_FONT
        cell.alignment = LEFT
    ws1.cell(row=row, column=1, value=f"ny-enhed-{i+1}")
    ws1.cell(row=row, column=15, value="tilføj")

# Column widths
col_widths = [22, 28, 12, 55, 10, 20, 35, 14, 14, 12, 12, 35, 45, 35, 22, 30]
for i, w in enumerate(col_widths):
    ws1.column_dimensions[get_column_letter(i + 1)].width = w

ws1.freeze_panes = "C2"


# ══════════════════════════════════════════════════════════════
# Sheet 2: USE CASES
# ══════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Use Cases")

uc_headers = ["ID", "Ikon", "Titel", "Beskrivelse", "HANDLING (behold/fjern/opdater)", "NOTER"]

for col, header in enumerate(uc_headers, 1):
    ws2.cell(row=1, column=col, value=header)
style_header(ws2, 1, len(uc_headers))

use_cases = [
    ["lighting", "💡", "Lysstyring", "Tænd, sluk og automatiser belysning i dit hjem eller kontor."],
    ["garage", "🚗", "Garageportstyring", "Åbn og luk din garageport eller port fra din telefon."],
    ["blinds", "🪟", "Persienner & Gardiner", "Automatiser rullegardiner, persienner og markiser."],
    ["energy", "⚡", "Energimåling", "Overvåg strømforbrug på individuelle kredse eller hele installationen."],
    ["dimming", "🔆", "Dæmpning", "Dæmp belysning jævnt med kompatible lyskilder og drivere."],
    ["appliance", "🔌", "Apparatstyring", "Gør enhver stikkontakt eller apparat smart med fjernbetjening."],
    ["heating", "🌡️", "Varme & Klima", "Styr radiatorer, gulvvarme eller klimaanlæg intelligent."],
    ["scenes", "🎬", "Scenestyring", "Konfigurer knapper til at aktivere scener og automatiseringer."],
    ["sensors", "📡", "Sensorovervågning", "Overvåg temperatur, fugtighed, oversvømmelse og bevægelse."],
    ["intercom", "🔔", "Dørtelefon & Adgang", "Automatiser dørtelefoner, låse og adgangskontrol."],
    ["motor", "⚙️", "Motorstyring", "Styr pumper, ventiler og andre motordrevne enheder."],
    ["display", "🖥️", "Centralt kontrolpanel", "Touchskærm til styring af alle enheder fra ét sted."],
]

for i, uc in enumerate(use_cases):
    row = i + 2
    for col, val in enumerate(uc, 1):
        ws2.cell(row=row, column=col, value=val)
    fill = GRAY_FILL if i % 2 == 0 else WHITE_FILL
    style_row(ws2, row, len(uc_headers), fill)
    ws2.cell(row=row, column=5).fill = YELLOW_FILL
    ws2.cell(row=row, column=6).fill = YELLOW_FILL

# Empty rows for new use cases
for i in range(5):
    row = len(use_cases) + 2 + i
    for col in range(1, len(uc_headers) + 1):
        cell = ws2.cell(row=row, column=col)
        cell.fill = GREEN_FILL
        cell.border = THIN_BORDER
    ws2.cell(row=row, column=1, value=f"ny-usecase-{i+1}")
    ws2.cell(row=row, column=5, value="tilføj")

uc_widths = [18, 6, 24, 60, 22, 30]
for i, w in enumerate(uc_widths):
    ws2.column_dimensions[get_column_letter(i + 1)].width = w

ws2.freeze_panes = "A2"


# ══════════════════════════════════════════════════════════════
# Sheet 3: FILTRE
# ══════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Filtre")

f_headers = ["Filter-gruppe", "Ikon", "Titel", "Beskrivelse",
             "Mulighed ID", "Mulighed Label", "Mulighed Værdi",
             "HANDLING (behold/fjern/opdater)", "NOTER"]

for col, header in enumerate(f_headers, 1):
    ws3.cell(row=1, column=col, value=header)
style_header(ws3, 1, len(f_headers))

filters = [
    ["voltage", "⚡", "Spændingsniveau", "Hvilken forsyningsspænding har du?",
     "any-v", "Alle", "any"],
    ["voltage", "", "", "", "230v", "110-240V AC", "110-240V AC"],
    ["voltage", "", "", "", "24vdc", "24V DC", "24V DC"],
    ["voltage", "", "", "", "12vdc", "12V DC", "12V DC"],
    ["voltage", "", "", "", "5vdc", "5V DC", "5-24V DC"],
    ["voltage", "", "", "", "battery", "Batteri", "Batteri"],

    ["installation", "🔧", "Installationsmetode", "Hvor skal enheden monteres?",
     "any-inst", "Alle", "any"],
    ["installation", "", "", "", "flush", "Indbygning (bag kontakt)", "flush"],
    ["installation", "", "", "", "din", "Tavlemontering (DIN-skinne)", "din"],
    ["installation", "", "", "", "plug", "Stikkontakt (Plug & Play)", "plug"],
    ["installation", "", "", "", "standalone", "Fritplaceret", "standalone"],

    ["powerMonitoring", "📊", "Effektmåling", "Skal enheden kunne måle strømforbrug?",
     "any-pm", "Ligegyldigt", "any"],
    ["powerMonitoring", "", "", "", "pm-yes", "Ja, krævet", "true"],
    ["powerMonitoring", "", "", "", "pm-no", "Nej, ikke nødvendigt", "false"],

    ["channels", "🔢", "Antal kanaler", "Hvor mange enheder skal styres?",
     "any-ch", "Alle", "any"],
    ["channels", "", "", "", "ch-1", "1 kanal", "1"],
    ["channels", "", "", "", "ch-2", "2 kanaler", "2"],
    ["channels", "", "", "", "ch-3", "3+ kanaler", "3"],

    ["dryContact", "🔗", "Tørkontakt", "Har du brug for potentialfrit relæ?",
     "any-dc", "Ligegyldigt", "any"],
    ["dryContact", "", "", "", "dc-yes", "Ja, krævet", "true"],
    ["dryContact", "", "", "", "dc-no", "Nej", "false"],

    ["protocol", "📡", "Protokol", "Hvilke protokoller er vigtige for dig?",
     "any-proto", "Alle", "any"],
    ["protocol", "", "", "", "proto-matter", "Matter", "Matter"],
    ["protocol", "", "", "", "proto-zigbee", "Zigbee", "Zigbee"],
    ["protocol", "", "", "", "proto-lan", "LAN (Ethernet)", "LAN"],
    ["protocol", "", "", "", "proto-wifi", "WiFi", "WiFi"],
]

for i, f in enumerate(filters):
    row = i + 2
    for col, val in enumerate(f, 1):
        ws3.cell(row=row, column=col, value=val)
    # Group header rows get light blue
    is_group_header = f[1] != ""
    fill = LIGHT_BLUE_FILL if is_group_header else (GRAY_FILL if i % 2 == 0 else WHITE_FILL)
    style_row(ws3, row, len(f_headers), fill)
    ws3.cell(row=row, column=8).fill = YELLOW_FILL
    ws3.cell(row=row, column=9).fill = YELLOW_FILL

f_widths = [18, 6, 22, 40, 16, 28, 18, 22, 30]
for i, w in enumerate(f_widths):
    ws3.column_dimensions[get_column_letter(i + 1)].width = w

ws3.freeze_panes = "A2"


# ══════════════════════════════════════════════════════════════
# Sheet 4: VEJLEDNING
# ══════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("Vejledning")

instructions = [
    ["VEJLEDNING TIL REDIGERING", ""],
    ["", ""],
    ["Ark: Enheder", "Hovedliste over alle Shelly-enheder med specifikationer og filtre."],
    ["  - Gule kolonner", "Udfyld 'HANDLING' med: behold, fjern, eller opdater"],
    ["  - Grønne rækker", "Tomme rækker til at tilføje nye enheder"],
    ["  - Use Cases kolonne", "Kommaseparerede use case ID'er (se 'Use Cases' arket)"],
    ["  - NOTER kolonne", "Skriv eventuelle rettelser eller kommentarer her"],
    ["", ""],
    ["Ark: Use Cases", "Liste over alle anvendelseskategorier i vælgeren."],
    ["  - Tilføj/fjern", "Brug HANDLING-kolonnen til at angive ændringer"],
    ["  - Nye use cases", "Udfyld de grønne rækker med nye kategorier"],
    ["", ""],
    ["Ark: Filtre", "Alle filtergrupper og deres valgmuligheder."],
    ["  - Blå rækker", "Filtergruppe-overskrifter"],
    ["  - Hvide rækker", "Individuelle filtermuligheder"],
    ["  - Tilføj nye filtre", "Skriv dem i NOTER-kolonnen"],
    ["", ""],
    ["Gyldige Use Case ID'er:", "lighting, garage, blinds, energy, dimming, appliance, heating, scenes, sensors, intercom, motor, display"],
    ["Gyldige Installations-typer:", "flush (indbygning), din (DIN-skinne), plug (stikkontakt), standalone (fritplaceret)"],
    ["Gyldige Protokoller:", "WiFi, Bluetooth, Zigbee, Matter, LAN"],
]

for i, (a, b) in enumerate(instructions):
    ws4.cell(row=i + 1, column=1, value=a)
    ws4.cell(row=i + 1, column=2, value=b)
    if i == 0:
        ws4.cell(row=i + 1, column=1).font = Font(name="Calibri", bold=True, size=14, color="0D1B2A")
    elif a.startswith("Ark:"):
        ws4.cell(row=i + 1, column=1).font = Font(name="Calibri", bold=True, size=11, color="0066FF")
        ws4.cell(row=i + 1, column=2).font = NORMAL_FONT
    elif a.startswith("  -"):
        ws4.cell(row=i + 1, column=1).font = Font(name="Calibri", bold=True, size=10, color="5A6577")
        ws4.cell(row=i + 1, column=2).font = NORMAL_FONT
    elif a.startswith("Gyldige"):
        ws4.cell(row=i + 1, column=1).font = Font(name="Calibri", bold=True, size=10, color="0D1B2A")
        ws4.cell(row=i + 1, column=2).font = Font(name="Calibri", size=10, color="0066FF")
    else:
        ws4.cell(row=i + 1, column=1).font = NORMAL_FONT
        ws4.cell(row=i + 1, column=2).font = NORMAL_FONT

ws4.column_dimensions["A"].width = 30
ws4.column_dimensions["B"].width = 80


# ─── Save ────────────────────────────────────────────────────
output_path = "/home/user/hello-world/shelly_enheder_oversigt.xlsx"
wb.save(output_path)
print(f"Excel file saved: {output_path}")
